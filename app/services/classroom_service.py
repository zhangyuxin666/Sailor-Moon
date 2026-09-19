import uuid
import mimetypes
from io import BytesIO
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from ..core.auth import create_account, require_manager
from ..core.audit import write_audit
from ..core.exceptions import NotFoundError, PermissionDeniedError
from ..config import settings
from ..integrations.qq import QQBotService
from ..models.schemas import RosterMemberIn
from .storage_service import StorageService

ALLOWED_UPLOAD_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".txt", ".zip", ".png", ".jpg", ".jpeg",
}


class ClassroomService:
    def __init__(self, db, reminders=None):
        self.db = db
        self.reminders = reminders
        self.qq = QQBotService(db)
        self.storage = StorageService()

    def _account_organization(self, conn, account_id: str) -> str:
        row = conn.execute(
            "SELECT organization_id FROM organization_members WHERE account_id = ? "
            "ORDER BY created_at LIMIT 1",
            (account_id,),
        ).fetchone()
        if not row:
            account = conn.execute(
                "SELECT username, role FROM accounts WHERE id = ?", (account_id,)
            ).fetchone()
            if not account or account["role"] != "manager":
                raise PermissionDeniedError("账号尚未加入组织")
            organization_id = uuid.uuid4().hex[:12]
            now = datetime.now().isoformat()
            conn.execute(
                "INSERT INTO organizations (id, name, owner_account_id, created_at) VALUES (?, ?, ?, ?)",
                (organization_id, f"{account['username']} 的组织", account_id, now),
            )
            conn.execute(
                "INSERT INTO organization_members (organization_id, account_id, role, created_at) "
                "VALUES (?, ?, 'owner', ?)",
                (organization_id, account_id, now),
            )
            return organization_id
        return row["organization_id"]

    def create_class(self, account: dict, name: str) -> dict:
        require_manager(account)
        class_id = uuid.uuid4().hex[:12]
        now = datetime.now().isoformat()
        with self.db.connect() as conn:
            organization_id = self._account_organization(conn, account["id"])
            binding = conn.execute(
                "SELECT group_openid FROM qq_bindings WHERE user_id = ? AND status = 'active' "
                "ORDER BY created_at LIMIT 1",
                (account["username"],),
            ).fetchone()
            if not binding:
                bindings = conn.execute(
                    "SELECT group_openid FROM qq_bindings WHERE status = 'active' ORDER BY created_at"
                ).fetchall()
                binding = bindings[0] if len(bindings) == 1 else None
            conn.execute(
                "INSERT INTO classes (id, name, manager_id, qq_group_openid, created_at) VALUES (?, ?, ?, ?, ?)",
                (class_id, name.strip(), account["id"], binding["group_openid"] if binding else None, now),
            )
            conn.execute(
                "INSERT INTO class_organizations (class_id, organization_id) VALUES (?, ?)",
                (class_id, organization_id),
            )
            write_audit(
                conn, "class.create", account["id"], organization_id,
                "class", class_id, {"name": name.strip()},
            )
        return {"id": class_id, "name": name, "qq_group_bound": bool(binding)}

    def _require_class_manager(self, conn, class_id: str, account: dict):
        row = conn.execute("SELECT * FROM classes WHERE id = ?", (class_id,)).fetchone()
        if not row:
            raise NotFoundError("班级不存在")
        organization = conn.execute(
            "SELECT co.organization_id, om.role FROM class_organizations co "
            "JOIN organization_members om ON om.organization_id = co.organization_id "
            "WHERE co.class_id = ? AND om.account_id = ?",
            (class_id, account["id"]),
        ).fetchone()
        if not organization or organization["role"] not in {"owner", "manager"} or account["role"] != "manager":
            raise PermissionDeniedError("只能管理自己创建的班级")
        result = dict(row)
        result["organization_id"] = organization["organization_id"]
        return result

    def import_roster(self, account: dict, class_id: str, members: list) -> dict:
        require_manager(account)
        created = []
        with self.db.connect() as conn:
            classroom = self._require_class_manager(conn, class_id, account)
            for item in members:
                existing = conn.execute(
                    "SELECT * FROM accounts WHERE username = ?", (item.username.strip(),)
                ).fetchone()
                if existing:
                    member = dict(existing)
                    if member["role"] != "participant":
                        raise ValueError(f"登录名 {item.username} 已被管理者占用")
                else:
                    member = create_account(
                        conn, item.username, item.display_name, item.password,
                        "participant", item.student_no,
                    )
                    conn.execute(
                        "INSERT INTO account_security (account_id, force_password_change, failed_logins, password_changed_at) "
                        "VALUES (?, 1, 0, ?)",
                        (member["id"], datetime.now().isoformat()),
                    )
                if not conn.execute(
                    "SELECT 1 FROM organization_members WHERE organization_id = ? AND account_id = ?",
                    (classroom["organization_id"], member["id"]),
                ).fetchone():
                    conn.execute(
                        "INSERT INTO organization_members (organization_id, account_id, role, created_at) "
                        "VALUES (?, ?, 'participant', ?)",
                        (classroom["organization_id"], member["id"], datetime.now().isoformat()),
                    )
                if not conn.execute(
                    "SELECT 1 FROM class_members WHERE class_id = ? AND account_id = ?",
                    (class_id, member["id"]),
                ).fetchone():
                    conn.execute(
                        "INSERT INTO class_members (class_id, account_id, created_at) VALUES (?, ?, ?)",
                        (class_id, member["id"], datetime.now().isoformat()),
                    )
                created.append({
                    "id": member["id"], "username": member["username"],
                    "display_name": member["display_name"], "student_no": item.student_no,
                })
            write_audit(
                conn, "roster.import", account["id"], classroom["organization_id"],
                "class", class_id, {"count": len(created)},
            )
        return {"count": len(created), "members": created}

    def import_roster_excel(self, account: dict, class_id: str, content: bytes) -> dict:
        if len(content) > 5 * 1024 * 1024:
            raise ValueError("Excel 文件不能超过 5MB")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
        except Exception as exc:
            raise ValueError("无法读取 Excel 文件，请使用 .xlsx 格式") from exc
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("Excel 文件为空")
        normalized = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
        required = ["学号", "姓名", "登录名", "初始密码"]
        missing_headers = [header for header in required if header not in normalized]
        if missing_headers:
            raise ValueError(f"缺少必填列：{'、'.join(missing_headers)}")

        def cell_text(value) -> str:
            if value is None:
                return ""
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()

        members = []
        usernames = set()
        for row_number, row in enumerate(rows, start=2):
            values = {header: cell_text(row[index] if index < len(row) else None) for header, index in normalized.items()}
            if not any(values.get(header) for header in required):
                continue
            if not all(values.get(header) for header in required):
                raise ValueError(f"第 {row_number} 行存在空白必填项")
            if values["登录名"] in usernames:
                raise ValueError(f"第 {row_number} 行登录名重复：{values['登录名']}")
            usernames.add(values["登录名"])
            members.append(RosterMemberIn(
                student_no=values["学号"], display_name=values["姓名"],
                username=values["登录名"], password=values["初始密码"],
            ))
            if len(members) > 500:
                raise ValueError("单次最多导入 500 名成员")
        if not members:
            raise ValueError("Excel 中没有可导入的成员数据")
        return self.import_roster(account, class_id, members)

    def create_todo(self, account: dict, class_id: str, data) -> dict:
        require_manager(account)
        try:
            deadline = datetime.fromisoformat(data.deadline)
        except ValueError as exc:
            raise ValueError("截止时间格式无效") from exc
        todo_id = uuid.uuid4().hex[:12]
        now = datetime.now().isoformat()
        with self.db.connect() as conn:
            classroom = self._require_class_manager(conn, class_id, account)
            conn.execute(
                "INSERT INTO todos (id, class_id, creator_id, title, description, kind, deadline, status, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, 'open', ?)",
                (todo_id, class_id, account["id"], data.title, data.description, data.kind, deadline.isoformat(), now),
            )
            members = conn.execute(
                "SELECT account_id FROM class_members WHERE class_id = ?", (class_id,)
            ).fetchall()
            for member in members:
                conn.execute(
                    "INSERT INTO todo_assignees (todo_id, account_id, status, updated_at) VALUES (?, ?, 'pending', ?)",
                    (todo_id, member["account_id"], now),
                )
            write_audit(
                conn, "todo.create", account["id"], classroom["organization_id"],
                "todo", todo_id, {"title": data.title, "assigned_count": len(members)},
            )
        if classroom["qq_group_openid"]:
            content = (
                f"【新{('作业' if data.kind == 'homework' else '待办')}】{data.title}\n"
                f"截止时间：{deadline.strftime('%Y-%m-%d %H:%M')}\n{data.description}\n"
                f"请登录活动管家平台完成提交。"
            )
            self._send_and_record(todo_id, classroom["qq_group_openid"], content, "publish")
        return {"id": todo_id, "title": data.title, "assigned_count": len(members)}

    def dashboard(self, account: dict) -> dict:
        with self.db.connect() as conn:
            if account["role"] == "manager":
                classes = [dict(row) for row in conn.execute(
                    "SELECT c.* FROM classes c JOIN class_organizations co ON co.class_id = c.id "
                    "JOIN organization_members om ON om.organization_id = co.organization_id "
                    "WHERE om.account_id = ? AND om.role IN ('owner', 'manager') ORDER BY c.created_at DESC",
                    (account["id"],),
                ).fetchall()]
                for item in classes:
                    item["member_count"] = conn.execute(
                        "SELECT COUNT(*) AS count FROM class_members WHERE class_id = ?",
                        (item["id"],),
                    ).fetchone()["count"]
                class_ids = [item["id"] for item in classes]
                todos_result = []
                for class_id in class_ids:
                    rows = conn.execute(
                        "SELECT t.*, c.name AS class_name FROM todos t JOIN classes c ON c.id = t.class_id "
                        "WHERE t.class_id = ? ORDER BY t.created_at DESC", (class_id,)
                    ).fetchall()
                    for row in rows:
                        item = dict(row)
                        counts = conn.execute(
                            "SELECT status, COUNT(*) AS count FROM todo_assignees WHERE todo_id = ? GROUP BY status",
                            (item["id"],),
                        ).fetchall()
                        summary = {entry["status"]: entry["count"] for entry in counts}
                        item["done_count"] = summary.get("done", 0)
                        item["total_count"] = sum(summary.values())
                        todos_result.append(item)
                return {"role": "manager", "classes": classes, "todos": todos_result}
            rows = conn.execute(
                "SELECT t.*, c.name AS class_name, ta.status AS my_status, ta.note, ta.original_filename, ta.submitted_at "
                "FROM todo_assignees ta JOIN todos t ON t.id = ta.todo_id "
                "JOIN classes c ON c.id = t.class_id WHERE ta.account_id = ? "
                "ORDER BY t.deadline",
                (account["id"],),
            ).fetchall()
            activity_rows = conn.execute(
                "SELECT t.id, t.title, t.status AS my_status, a.title AS activity_title, "
                "a.plan_json, a.created_at FROM activity_task_assignees ata "
                "JOIN tasks t ON t.id = ata.task_id JOIN activities a ON a.id = t.activity_id "
                "WHERE ata.account_id = ? ORDER BY a.created_at DESC",
                (account["id"],),
            ).fetchall()
            import json

            activity_tasks = []
            for row in activity_rows:
                item = dict(row)
                plan = json.loads(item.pop("plan_json")) if item.get("plan_json") else {}
                item["event_time"] = plan.get("event_time") or item["created_at"]
                activity_tasks.append(item)
            return {
                "role": "participant",
                "todos": [dict(row) for row in rows],
                "activity_tasks": activity_tasks,
            }

    def todo_detail(self, account: dict, todo_id: str) -> dict:
        with self.db.connect() as conn:
            todo = conn.execute(
                "SELECT t.*, c.name AS class_name, c.manager_id FROM todos t "
                "JOIN classes c ON c.id = t.class_id WHERE t.id = ?", (todo_id,)
            ).fetchone()
            if not todo:
                raise NotFoundError("待办不存在")
            if account["role"] == "manager":
                classroom = self._require_class_manager(conn, todo["class_id"], account)
                rows = conn.execute(
                    "SELECT a.id AS account_id, a.display_name, a.student_no, a.username, "
                    "ta.status, ta.note, ta.original_filename, ta.submitted_at "
                    "FROM todo_assignees ta JOIN accounts a ON a.id = ta.account_id "
                    "WHERE ta.todo_id = ? ORDER BY a.student_no, a.display_name",
                    (todo_id,),
                ).fetchall()
                reminders = conn.execute(
                    "SELECT * FROM todo_reminders WHERE todo_id = ? ORDER BY remind_at", (todo_id,)
                ).fetchall()
                deliveries = conn.execute(
                    "SELECT * FROM classroom_deliveries WHERE todo_id = ? ORDER BY created_at DESC",
                    (todo_id,),
                ).fetchall()
                return {
                    "todo": dict(todo),
                    "participants": [dict(row) for row in rows],
                    "reminders": [dict(row) for row in reminders],
                    "deliveries": [dict(row) for row in deliveries],
                }
            assignment = conn.execute(
                "SELECT * FROM todo_assignees WHERE todo_id = ? AND account_id = ?",
                (todo_id, account["id"]),
            ).fetchone()
            if not assignment:
                raise PermissionDeniedError("该待办未分配给你")
            return {"todo": dict(todo), "assignment": dict(assignment)}

    def submit(self, account: dict, todo_id: str, note: str, filename: str | None, content: bytes | None) -> dict:
        if account["role"] != "participant":
            raise PermissionDeniedError("只有参与者可以提交待办")
        with self.db.connect() as conn:
            assigned = conn.execute(
                "SELECT 1 FROM todo_assignees WHERE todo_id = ? AND account_id = ?",
                (todo_id, account["id"]),
            ).fetchone()
        if not assigned:
            raise PermissionDeniedError("该待办未分配给你")
        stored_path = None
        safe_name = None
        if content is not None and filename:
            if len(content) > settings.upload_max_mb * 1024 * 1024:
                raise ValueError(f"文件不能超过 {settings.upload_max_mb}MB")
            safe_name = Path(filename).name[:180]
            extension = Path(safe_name).suffix.lower()
            if extension not in ALLOWED_UPLOAD_EXTENSIONS:
                raise ValueError("不支持该文件类型")
            with self.db.connect() as conn:
                used = conn.execute(
                    "SELECT COALESCE(SUM(size_bytes), 0) AS total FROM file_assets WHERE account_id = ?",
                    (account["id"],),
                ).fetchone()["total"]
            quota = settings.user_storage_quota_mb * 1024 * 1024
            if used + len(content) > quota:
                raise ValueError(f"个人存储空间不能超过 {settings.user_storage_quota_mb}MB")
            stored_path = self.storage.save(content, safe_name, f"todos/{todo_id}/{account['id']}")
        now = datetime.now().isoformat()
        with self.db.connect() as conn:
            if stored_path:
                asset_id = uuid.uuid4().hex[:16]
                conn.execute(
                    "UPDATE todo_assignees SET status = 'done', note = ?, file_path = ?, original_filename = ?, submitted_at = ?, updated_at = ? "
                    "WHERE todo_id = ? AND account_id = ?",
                    (note, asset_id, safe_name, now, now, todo_id, account["id"]),
                )
                conn.execute(
                    "INSERT INTO file_assets (id, todo_id, account_id, storage_key, original_filename, content_type, size_bytes, created_at) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        asset_id, todo_id, account["id"], stored_path, safe_name,
                        mimetypes.guess_type(safe_name)[0] or "application/octet-stream",
                        len(content), now,
                    ),
                )
            else:
                conn.execute(
                    "UPDATE todo_assignees SET status = 'done', note = ?, submitted_at = ?, updated_at = ? "
                    "WHERE todo_id = ? AND account_id = ?",
                    (note, now, now, todo_id, account["id"]),
                )
            organization_id = conn.execute(
                "SELECT co.organization_id FROM todos t JOIN class_organizations co ON co.class_id = t.class_id "
                "WHERE t.id = ?", (todo_id,),
            ).fetchone()["organization_id"]
            write_audit(
                conn, "todo.submit", account["id"], organization_id,
                "todo", todo_id, {"has_file": bool(stored_path)},
            )
        return {"todo_id": todo_id, "status": "done", "submitted_at": now}

    def schedule_reminder(self, account: dict, todo_id: str, remind_at_text: str) -> dict:
        require_manager(account)
        remind_at = datetime.fromisoformat(remind_at_text)
        with self.db.connect() as conn:
            todo = conn.execute("SELECT class_id FROM todos WHERE id = ?", (todo_id,)).fetchone()
            if not todo:
                raise NotFoundError("待办不存在")
            self._require_class_manager(conn, todo["class_id"], account)
            return self.reminders.schedule_todo(conn, todo_id, remind_at)

    def remind_missing(self, account: dict | None, todo_id: str, scheduled: bool = False) -> dict:
        with self.db.connect() as conn:
            todo = conn.execute(
                "SELECT t.*, c.manager_id, c.qq_group_openid FROM todos t JOIN classes c ON c.id = t.class_id "
                "WHERE t.id = ?", (todo_id,)
            ).fetchone()
            if not todo:
                raise NotFoundError("待办不存在")
            if not scheduled:
                if not account:
                    raise PermissionDeniedError("无权催办该任务")
                self._require_class_manager(conn, todo["class_id"], account)
            missing = conn.execute(
                "SELECT a.display_name FROM todo_assignees ta JOIN accounts a ON a.id = ta.account_id "
                "WHERE ta.todo_id = ? AND ta.status != 'done' ORDER BY a.student_no, a.display_name",
                (todo_id,),
            ).fetchall()
        names = [row["display_name"] for row in missing]
        if not names:
            return {"status": "skipped", "message": "所有人都已完成", "missing": []}
        if not todo["qq_group_openid"]:
            raise NotFoundError("班级尚未绑定 QQ 群")
        content = (
            f"【未完成提醒】{todo['title']}\n截止时间：{todo['deadline']}\n"
            f"以下同学尚未完成：{'、'.join(names)}\n请尽快登录活动管家平台提交。"
        )
        result = self._send_and_record(todo_id, todo["qq_group_openid"], content, "scheduled" if scheduled else "manual")
        return {"status": "sent", "missing": names, "delivery": result}

    def _send_and_record(self, todo_id: str, group_openid: str, content: str, kind: str) -> dict:
        delivery_id = uuid.uuid4().hex[:12]
        try:
            result = self.qq.send_group(group_openid, content)
            status, error, external_id = "sent", None, result.get("id")
        except Exception as exc:
            status, error, external_id = "failed", str(exc)[:1000], None
            result = {"status": "failed", "error": error}
        with self.db.connect() as conn:
            conn.execute(
                "INSERT INTO classroom_deliveries (id, todo_id, kind, content, status, external_id, error, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (delivery_id, todo_id, kind, content, status, external_id, error, datetime.now().isoformat()),
            )
        if status == "failed":
            raise RuntimeError(error)
        return result
